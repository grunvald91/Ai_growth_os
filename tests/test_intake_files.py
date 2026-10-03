import asyncio
import io
import subprocess
import zipfile

import pytest
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.intake.sources import MAX_FILE_BYTES, MAX_TEXT, make_source, parse_file, parse_file_async


def pdf_bytes():
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                             NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 16 Tf 50 700 Td (Example company massage price 5000) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def image_bytes(text='EXAMPLE PRICE 5000'):
    image = Image.new('RGB', (1200, 160), 'white')
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 48)
    draw.text((20, 40), text, fill='black', font=font)
    output = io.BytesIO()
    image.save(output, format='PNG')
    return output.getvalue()


def test_pdf_real_text_and_encrypted_pdf():
    source = parse_file('company.pdf', pdf_bytes(), 'pdf')
    assert source.status == 'read' and 'massage price 5000' in source.text
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt('secret')
    output = io.BytesIO()
    writer.write(output)
    source = parse_file('protected.pdf', output.getvalue(), 'protected')
    assert source.status == 'unsupported' and not source.text


def test_spreadsheet_multiple_sheets_and_missing_formula_values():
    book = Workbook()
    book.active.title = 'Услуги'
    book.active.append(['Массаж', 5000])
    book.create_sheet('Контакты').append(['Компания Example'])
    output = io.BytesIO()
    book.save(output)
    source = parse_file('price.xlsx', output.getvalue(), 'xlsx')
    assert source.status == 'read' and '5000' in source.text and 'Компания Example' in source.text
    book.active.append(['Итог', '=B1*2'])
    output = io.BytesIO()
    book.save(output)
    source = parse_file('price.xlsx', output.getvalue(), 'formula')
    assert source.status == 'truncated' and 'без сохранённого результата' in source.text
    assert 'Формулы не вычисляются' in source.note
    assert '10000' not in source.text


def test_empty_and_partial_files_are_not_claimed_read():
    for data in [b'', b'  \n\t']:
        source = parse_file('empty.txt', data, 'empty')
        assert source.status == 'unavailable' and source.note and not source.text
    source = parse_file('long.txt', b'a' * (MAX_TEXT + 100), 'long')
    assert source.status == 'truncated' and len(source.text) == MAX_TEXT and source.note
    assert make_source('owner_message', 'message', 'x' * 13000).note
    assert parse_file('large.txt', b'x' * (MAX_FILE_BYTES + 1), 'large').status == 'unsupported'


def test_semicolon_csv_and_russian_encoding():
    source = parse_file('price.csv', 'Услуга;Цена\nМассаж;5000'.encode('cp1251'), 'csv')
    assert source.status == 'read' and 'Массаж | 5000' in source.text
    source = parse_file('company.txt', 'Описание компании'.encode('utf-16'), 'txt')
    assert source.status == 'read' and 'Описание' in source.text


def test_office_entities_and_macros_rejected():
    for xml in [b'<!DOCTYPE x [<!ENTITY e "boom">]><x>&e;</x>', b'<broken>']:
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w') as archive:
            archive.writestr('word/document.xml', xml)
        assert parse_file('company.docx', output.getvalue(), 'xml').status == 'unsupported'
    assert parse_file('company.xlsm', b'macro', 'macro').status == 'unsupported'


def test_ocr_image_and_scanned_pdf():
    langs = subprocess.run(['tesseract', '--list-langs'], capture_output=True, text=True, check=True).stdout
    if 'rus' not in langs:
        pytest.skip('Russian OCR data is required in CI and production')
    source = parse_file('scan.png', image_bytes('МАССАЖ ЦЕНА 5000'), 'image')
    assert source.status == 'read' and '5000' in source.text and 'МАССАЖ' in source.text
    image = Image.open(io.BytesIO(image_bytes()))
    output = io.BytesIO()
    image.save(output, format='PDF')
    source = parse_file('scan.pdf', output.getvalue(), 'scan')
    assert source.status == 'read' and '5000' in source.text and 'распознан' in source.note


@pytest.mark.asyncio
async def test_isolated_parser_and_timeout(monkeypatch):
    source = await parse_file_async('company.pdf', pdf_bytes(), 'pdf')
    assert source.status == 'read' and '5000' in source.text
    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired('parser', 90)
    monkeypatch.setattr('app.intake.sources.subprocess.run', timeout)
    source = await parse_file_async('company.pdf', b'fake', 'timeout')
    assert source.status == 'unavailable' and source.note and not source.text


@pytest.mark.asyncio
async def test_parser_does_not_block_event_loop(monkeypatch):
    import time
    def slow(*args, **kwargs):
        time.sleep(.15)
        return subprocess.CompletedProcess(args, 0, stdout=make_source('owner_file', 'test', 'Company').model_dump_json().encode())
    monkeypatch.setattr('app.intake.sources.subprocess.run', slow)
    parsing = asyncio.create_task(parse_file_async('company.txt', b'Company', 'test'))
    await asyncio.sleep(.02)
    assert not parsing.done()
    assert (await parsing).status == 'read'


@pytest.mark.asyncio
async def test_legacy_excel_real_file_isolated():
    from pathlib import Path
    data = (Path(__file__).parent / 'fixtures/intake-price.xls').read_bytes()
    source = await parse_file_async('price.xls', data, 'xls')
    assert source.status == 'read' and 'Массаж | 5000' in source.text


@pytest.mark.asyncio
async def test_xlsx_isolated():
    book = Workbook()
    book.active.append(['Массаж', 5000])
    output = io.BytesIO()
    book.save(output)
    source = await parse_file_async('price.xlsx', output.getvalue(), 'xlsx-process')
    assert source.status == 'read' and 'Массаж | 5000' in source.text


@pytest.mark.asyncio
async def test_long_filename_keeps_format_extension():
    source = await parse_file_async('company-' + 'x' * 210 + '.csv', b'service;price\nMassage;5000', 'long-name')
    assert source.status == 'read' and 'Massage | 5000' in source.text
    assert len(source.filename) == 200
