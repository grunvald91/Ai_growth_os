"""Local document extraction. Invoked in a bounded subprocess by async intake."""
import io
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree

LIMIT = 12_001


def office_archive(data):
    archive = zipfile.ZipFile(io.BytesIO(data))
    infos = archive.infolist()
    if (len(infos) > 2000 or sum(i.file_size for i in infos) > 64 * 1024 * 1024
            or any(i.file_size > 16 * 1024 * 1024 or "vba" in i.filename.casefold() for i in infos)):
        archive.close()
        raise ValueError("Unsafe office container")
    return archive


def xml_root(data):
    if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
        raise ValueError("Unsafe XML")
    return ElementTree.fromstring(data)


def ocr_image(data):
    from PIL import Image
    with Image.open(io.BytesIO(data)) as image:
        if image.width * image.height > 25_000_000:
            raise ValueError("Image too large")
        image.load()
        with tempfile.TemporaryDirectory(prefix="intake-image-") as folder:
            path = Path(folder) / "image.png"
            image.convert("RGB").save(path)
            result = subprocess.run(["tesseract", str(path), "stdout", "-l", "rus+eng"],
                                    capture_output=True, timeout=20, check=True)
            return result.stdout.decode("utf-8")


def extract_pdf(data):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ValueError("Encrypted PDF")
    parts, notes = [], []
    with tempfile.TemporaryDirectory(prefix="intake-pdf-") as folder:
        path = Path(folder) / "document.pdf"
        path.write_bytes(data)
        for index, page in enumerate(reader.pages[:40]):
            text = page.extract_text() or ""
            if not text.strip():
                prefix = Path(folder) / "page"
                subprocess.run(["pdftoppm", "-f", str(index + 1), "-l", str(index + 1),
                                "-scale-to", "2400", "-singlefile", "-png", str(path), str(prefix)],
                               capture_output=True, timeout=15, check=True)
                text = ocr_image(prefix.with_suffix(".png").read_bytes())
                notes.append("Текст на сканах распознан автоматически; проверьте названия, числа и цены.")
            parts.append(f"Страница {index + 1}\n{text}")
            if sum(len(p) for p in parts) > LIMIT:
                break
        partial = len(parts) < len(reader.pages)
    if partial:
        notes.append(f"Прочитано страниц: {len(parts)} из {len(reader.pages)}. Пришлите нужные страницы отдельным файлом.")
    # Page markers alone must not make an empty PDF look successfully read.
    meaningful = any(part.split("\n", 1)[1].strip() for part in parts)
    return "\n\n".join(parts) if meaningful else "", partial, " ".join(dict.fromkeys(notes))


def extract_spreadsheet(suffix, data):
    parts, length, missing_formulas = [], 0, False
    partial = False
    books = []
    try:
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            with office_archive(data):
                pass
            values = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            books.append(values)
            formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            books.append(formulas)
            sheets = ((sheet.title, sheet.iter_rows(), formulas[sheet.title].iter_rows()) for sheet in values)
        else:
            import xlrd
            book = xlrd.open_workbook(file_contents=data, on_demand=True)
            books.append(book)
            sheets = ((s.name, ([s.cell_value(r, c) for c in range(s.ncols)] for r in range(s.nrows)), None)
                      for s in book.sheets())
        cells = 0
        for title, rows, formula_rows in sheets:
            parts.append("Лист: " + title)
            length += len(parts[-1])
            if formula_rows is not None:
                rows = zip(rows, formula_rows)
            for row in rows:
                if formula_rows is not None:
                    values_row, formulas_row = row
                    values = []
                    for cell, formula in zip(values_row, formulas_row):
                        if formula.data_type == "f" and cell.value is None:
                            missing_formulas = True
                            values.append("[формула без сохранённого результата]")
                        else:
                            values.append("" if cell.value is None else str(cell.value))
                else:
                    values = [str(v) if v != "" else "" for v in row]
                cells += len(values)
                if cells > 100_000:
                    partial = True
                    break
                line = " | ".join(values)
                if line.strip(" |"):
                    parts.append(line)
                    length += len(line) + 1
                if length > LIMIT:
                    partial = True
                    break
            if partial:
                break
        note = "Формулы не вычисляются. Есть ячейки без сохранённого результата; сохраните вычисленные значения и загрузите файл снова." if missing_formulas else ""
        if partial:
            note += " Таблица прочитана частично; пришлите нужные листы или строки отдельно."
        # Empty sheets are not usable company evidence.
        text = "\n".join(parts) if any(not p.startswith("Лист: ") for p in parts) else ""
        return text, partial or missing_formulas, note.strip()
    finally:
        for book in books:
            close = getattr(book, "close", None) or getattr(book, "release_resources", None)
            if close:
                close()


if __name__ == "__main__":
    import resource
    resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024,) * 2)
    resource.setrlimit(resource.RLIMIT_CPU, (60, 60))
    from app.intake.sources import parse_file
    source = parse_file(sys.argv[1], sys.stdin.buffer.read(), sys.argv[2])
    sys.stdout.write(source.model_dump_json())
